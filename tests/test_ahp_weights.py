"""AHP weights: read from the sheet, computed in code, refused when inconsistent.

The committed sheet is checked against its known priority vector, and small
workbooks written on the fly pin the method (geometric mean, Saaty's CR) and
the refusal of judgments with CR >= 0.10.
"""
import math

import openpyxl
import pytest

from pipeline import ahp, scoring

KEYS = [
    "Market Signal Strength",
    "Evidence Quality",
    "Source Diversity",
    "Strategic Relevance",
    "Novelty Momentum",
]


def _sheet(tmp_path, upper, names=KEYS):
    """Write a workbook laid out like pipeline/ahp_weights.xlsx; `upper` maps (i, j) -> judgment."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for j, name in enumerate(names):
        ws.cell(ahp.HEADER_ROW, ahp.FIRST_COLUMN + j, name)
        ws.cell(ahp.HEADER_ROW + 1 + j, ahp.FIRST_COLUMN - 1, name)
        ws.cell(ahp.HEADER_ROW + 1 + j, ahp.FIRST_COLUMN + j, 1)
    for (i, j), value in upper.items():
        ws.cell(ahp.HEADER_ROW + 1 + i, ahp.FIRST_COLUMN + j, value)
    path = tmp_path / "ahp.xlsx"
    wb.save(path)
    return path


def test_committed_sheet_gives_known_weights_and_passes_consistency():
    weights, cr = ahp.ahp_weights()
    assert math.isclose(sum(weights.values()), 1.0)
    assert {k: round(v, 3) for k, v in weights.items()} == {
        "market_signal_strength": 0.397,
        "evidence_quality": 0.278,
        "source_diversity": 0.141,
        "strategic_relevance": 0.130,
        "novelty_momentum": 0.055,
    }
    assert round(cr, 3) == 0.021


def test_perfectly_consistent_judgments_recover_the_ratios(tmp_path):
    # weights 4:2:1 -> a/b = 2, a/c = 4, b/c = 2, which is exactly consistent.
    path = _sheet(tmp_path, {(0, 1): 2, (0, 2): 4, (1, 2): 2}, names=KEYS[:3])
    weights, cr = ahp.ahp_weights(path)
    assert [round(w, 6) for w in weights.values()] == [round(4 / 7, 6), round(2 / 7, 6), round(1 / 7, 6)]
    assert cr == pytest.approx(0.0, abs=1e-12)


def test_inconsistent_judgments_are_refused(tmp_path):
    # a >> b, b >> c, but c >> a: a cycle no weight vector can honour.
    path = _sheet(tmp_path, {(0, 1): 9, (0, 2): 1 / 9, (1, 2): 9}, names=KEYS[:3])
    with pytest.raises(ahp.InconsistentJudgments, match="consistency ratio"):
        ahp.ahp_weights(path)


def test_a_missing_or_non_positive_judgment_is_rejected(tmp_path):
    path = _sheet(tmp_path, {(0, 1): 2, (0, 2): 0, (1, 2): 2}, names=KEYS[:3])
    with pytest.raises(ValueError, match="positive number"):
        ahp.read_pairwise(path)


def test_lower_triangle_is_rebuilt_as_reciprocals(tmp_path):
    path = _sheet(tmp_path, {(0, 1): 3, (0, 2): 5, (1, 2): 2}, names=KEYS[:3])
    _, matrix = ahp.read_pairwise(path)
    assert matrix[1][0] == pytest.approx(1 / 3)
    assert matrix[2][0] == pytest.approx(1 / 5)
    assert matrix[2][1] == pytest.approx(1 / 2)


def test_manual_weighting_is_kept_as_a_named_option():
    assert scoring.attractiveness_weights("manual") == {
        "market_signal_strength": 0.3,
        "source_diversity": 0.2,
        "evidence_quality": 0.25,
        "novelty_momentum": 0.1,
        "strategic_relevance": 0.15,
    }


def test_unknown_weighting_is_refused():
    with pytest.raises(ValueError, match="Unknown weighting"):
        scoring.attractiveness_weights("equal")


def test_scoring_defaults_to_the_ahp_weights():
    assert scoring.WEIGHTING == "ahp"
    assert scoring.WEIGHTS == ahp.ahp_weights()[0]


def test_sheet_criteria_must_match_the_sub_scores(tmp_path, monkeypatch):
    path = _sheet(tmp_path, {(0, 1): 2}, names=["Market Signal Strength", "Hype"])
    monkeypatch.setattr(scoring, "ahp_weights", lambda: ahp.ahp_weights(path))
    with pytest.raises(ValueError, match="do not match the sub-scores"):
        scoring.attractiveness_weights("ahp")


# runpy warns that pipeline.scoring is already imported; re-running it is the point.
@pytest.mark.filterwarnings("ignore:.*found in sys.modules:RuntimeWarning")
def test_cli_refuses_an_unknown_weighting_before_scoring(monkeypatch):
    import runpy
    import sys

    monkeypatch.setattr(sys, "argv", ["scoring", "--weights=equal"])
    with pytest.raises(ValueError, match="Unknown weighting"):
        runpy.run_module("pipeline.scoring", run_name="__main__")
