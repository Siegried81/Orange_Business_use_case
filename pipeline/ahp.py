"""Attractiveness weights from an AHP pairwise-comparison sheet.

`pipeline/ahp_weights.xlsx` holds the judgments: one row and one column per
sub-score, and only the upper triangle is typed in (1 = equal ... 9 =
extreme). This module reads those typed cells, rebuilds the reciprocal lower
triangle itself, and computes the priority vector and its consistency ratio
in code, so the weights the scoring uses can never drift from a stale cached
value in the workbook.

The method is the one the sheet's own formulas use, so code and sheet agree:
the priority vector is the normalised geometric mean of each row, and
lambda_max is the mean of (A.w)_i / w_i. Judgments with CR >= 0.10 are
refused (Saaty's threshold): past it the comparisons contradict each other
too much for the vector to mean anything.
"""
import math
from pathlib import Path

import openpyxl

AHP_SHEET_PATH = Path(__file__).with_name("ahp_weights.xlsx")
CR_THRESHOLD = 0.10
# Saaty's random consistency index, by matrix size.
RANDOM_INDEX = {1: 0.0, 2: 0.0, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24, 7: 1.32, 8: 1.41, 9: 1.45}
# Where the matrix sits in the sheet: criterion names in row 6 from column B,
# one row per criterion from row 7.
HEADER_ROW = 6
FIRST_COLUMN = 2


class InconsistentJudgments(ValueError):
    """The pairwise judgments are too inconsistent (CR >= CR_THRESHOLD) to use."""


def _key(name):
    """'Market Signal Strength' -> 'market_signal_strength', the scoring key."""
    return "_".join(str(name).strip().lower().split())


def read_pairwise(path=AHP_SHEET_PATH):
    """Return (criteria keys, full reciprocal matrix) from the sheet's upper triangle.

    Only the typed upper-triangle cells are read; formula cells are ignored
    because openpyxl does not evaluate them.
    """
    ws = openpyxl.load_workbook(path).active
    names = []
    col = FIRST_COLUMN
    while ws.cell(HEADER_ROW, col).value:
        names.append(_key(ws.cell(HEADER_ROW, col).value))
        col += 1
    n = len(names)
    if n < 2:
        raise ValueError(f"{path}: fewer than two criteria in row {HEADER_ROW}")
    matrix = [[1.0] * n for _ in range(n)]
    for i in range(n):
        if _key(ws.cell(HEADER_ROW + 1 + i, FIRST_COLUMN - 1).value) != names[i]:
            raise ValueError(f"{path}: row {HEADER_ROW + 1 + i} does not match column {names[i]}")
        for j in range(i + 1, n):
            value = ws.cell(HEADER_ROW + 1 + i, FIRST_COLUMN + j).value
            if not isinstance(value, (int, float)) or value <= 0:
                raise ValueError(
                    f"{path}: judgment {names[i]} vs {names[j]} must be a positive number, got {value!r}"
                )
            matrix[i][j] = float(value)
            matrix[j][i] = 1.0 / float(value)
    return names, matrix


def priority_vector(matrix):
    """Normalised geometric mean of each row."""
    n = len(matrix)
    geo = [math.prod(row) ** (1.0 / n) for row in matrix]
    total = sum(geo)
    return [g / total for g in geo]


def consistency_ratio(matrix, weights):
    """CR = CI / RI with CI = (lambda_max - n) / (n - 1); 0.0 below 3 criteria."""
    n = len(matrix)
    if RANDOM_INDEX.get(n, 0.0) == 0.0:
        return 0.0
    lambdas = [sum(matrix[i][j] * weights[j] for j in range(n)) / weights[i] for i in range(n)]
    ci = (sum(lambdas) / n - n) / (n - 1)
    return ci / RANDOM_INDEX[n]


def ahp_weights(path=AHP_SHEET_PATH):
    """Return ({criterion: weight}, CR), refusing judgments with CR >= CR_THRESHOLD."""
    names, matrix = read_pairwise(path)
    weights = priority_vector(matrix)
    cr = consistency_ratio(matrix, weights)
    if cr >= CR_THRESHOLD:
        raise InconsistentJudgments(
            f"{path}: consistency ratio {cr:.3f} >= {CR_THRESHOLD}; revise the pairwise "
            "judgments before using them as weights."
        )
    return dict(zip(names, weights)), cr
